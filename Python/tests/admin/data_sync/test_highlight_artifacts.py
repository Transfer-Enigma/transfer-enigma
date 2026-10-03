from io import BytesIO

import pandas as pd
from backend_admin.service.routes_loading.artifacts import build_report_workbook
from backend_admin.service.routes_loading.sync_errors import SyncError, SyncErrorCode
from backend_admin.service.routes_loading.uid_sheet import highlight_report_cells, write_fix_cells
from openpyxl import load_workbook


class FakeSpreadsheet:
    def __init__(self):
        self.updates = []

    def batch_update(self, body):
        self.updates.append(body)


class FakeWorksheet:
    def __init__(self, values, sheet_id=7):
        self._values = [list(row) for row in values]
        self.id = sheet_id
        self.spreadsheet = FakeSpreadsheet()
        self.updates = []

    def get_all_values(self):
        return [list(row) for row in self._values]

    def update(self, cell_range, values):
        self.updates.append((cell_range, values))


def _finding(sheet="МОРЕ", row=2, cell="company"):
    return SyncError(
        document=None, sheet=sheet, row=row, cell=cell,
        code=SyncErrorCode.NO_PRICE, message="m",
    )


class TestWriteFixCells:
    def test_writes_grouped_by_sheet(self):
        ws = FakeWorksheet([["company", "city"], ["a", "b"]])
        fixes = [
            {"sheet": "sea", "row": 2, "column": "company", "old": "a", "new": "A"},
            {"sheet": "sea", "row": 2, "column": "unknown_col", "old": "x", "new": "y"},
            {"sheet": "other", "row": 2, "column": "company", "old": "a", "new": "A"},
        ]
        written = write_fix_cells({"ws": ws}, {"sea": ("SEA", "ws")}, fixes)

        assert written == 1
        assert ws.updates == [("A2", [["A"]])]

    def test_skips_missing_worksheets(self):
        assert write_fix_cells({}, {"sea": ("SEA", "ws")}, [
            {"sheet": "sea", "row": 2, "column": "c", "old": "a", "new": "b"},
        ]) == 0


class TestHighlightReportCells:
    def test_highlights_error_cells(self):
        ws = FakeWorksheet([["company", "city"], ["a", "b"]])
        findings = [_finding(), _finding(row=None), _finding(cell=None)]

        highlighted = highlight_report_cells({"ws": ws}, {"sea": ("SEA", "ws")}, findings)

        assert highlighted == 1
        request = ws.spreadsheet.updates[0]["requests"][0]["repeatCell"]
        assert request["range"] == {
            "sheetId": 7,
            "startRowIndex": 1,
            "endRowIndex": 2,
            "startColumnIndex": 0,
            "endColumnIndex": 1,
        }
        assert request["cell"]["userEnteredFormat"]["backgroundColor"] == {
            "red": 1.0, "green": 0.85, "blue": 0.85,
        }

    def test_unknown_columns_and_sheets_skipped(self):
        ws = FakeWorksheet([["company"], ["a"]])
        findings = [_finding(cell="nope"), _finding(sheet="ЖД")]

        assert highlight_report_cells({"ws": ws}, {"sea": ("SEA", "ws")}, findings) == 0
        assert ws.spreadsheet.updates == []


class TestBuildReportWorkbook:
    def _frames(self):
        return {
            "sea": pd.DataFrame([{"company": "a", "city": "b"}]),
            "rail": None,
            "truck": None,
            "dropp": None,
            "services": None,
            "points": None,
        }

    def test_error_and_fix_fills(self):
        content = build_report_workbook(
            self._frames(),
            [_finding()],
            [{"sheet": "sea", "row": 2, "column": "city", "old": "b", "new": "B"}],
            {"sea": "Море"},
        )
        workbook = load_workbook(BytesIO(content))
        sheet = workbook["Море"]

        assert sheet["A1"].value == "company"
        assert sheet["A2"].value == "a"
        assert "FFC7CE" in str(sheet["A2"].fill.start_color.rgb)
        assert "C6EFCE" in str(sheet["B2"].fill.start_color.rgb)

    def test_skips_empty_frames(self):
        content = build_report_workbook(self._frames(), [], [], {})
        workbook = load_workbook(BytesIO(content))

        assert workbook.sheetnames == ["sea"]
