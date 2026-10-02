from contextlib import suppress
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pandas import DataFrame

from .uid_sheet import SHEET_SCOPES

ERROR_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FIX_FILL = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")

SHEET_KEYS = ("sea", "rail", "truck", "dropp", "services", "points")


def _cell_value(value):
    with suppress(TypeError, ValueError):
        if pd.isna(value):
            return None
    return value


def _column_positions(frame: DataFrame) -> dict[str, int]:
    return {str(column): position for position, column in enumerate(frame.columns)}


def _scope_of(key: str) -> str:
    return {
        "sea": "SEA", "rail": "RAIL", "truck": "TRUCK", "dropp": "DROPP",
        "services": "services", "points": "points",
    }.get(key, key)


def _collect_error_cells(frames: dict[str, DataFrame | None], findings: list) -> dict[str, set]:
    error_cells: dict[str, set[tuple[int, int]]] = {}
    for finding in findings:
        scope = SHEET_SCOPES.get(getattr(finding, "sheet", None) or "")
        row = getattr(finding, "row", None)
        cell = getattr(finding, "cell", None)
        if scope is None or row is None or not cell:
            continue
        for key, frame in frames.items():
            if frame is None or _scope_of(key) != scope:
                continue
            positions = _column_positions(frame)
            for column in str(cell).split(","):
                column = column.strip()
                if column in positions:
                    error_cells.setdefault(key, set()).add((row, positions[column] + 1))
    return error_cells


def _collect_fix_cells(frames: dict[str, DataFrame | None], fixes: list[dict]) -> dict[str, set]:
    fix_cells: dict[str, set[tuple[int, int]]] = {}
    for fix in fixes:
        fix_key = fix.get("sheet")
        if not isinstance(fix_key, str):
            continue
        frame = frames.get(fix_key)
        if frame is None:
            continue
        positions = _column_positions(frame)
        if fix.get("column") in positions and fix.get("row"):
            fix_cells.setdefault(fix_key, set()).add((fix["row"], positions[fix["column"]] + 1))
    return fix_cells


def _paint_sheet(sheet, error_positions: set, fix_positions: set) -> None:
    for excel_row, excel_col in sorted(error_positions):
        if excel_row <= sheet.max_row:
            sheet.cell(row=excel_row, column=excel_col).fill = ERROR_FILL
    for excel_row, excel_col in sorted(fix_positions - error_positions):
        if excel_row <= sheet.max_row:
            sheet.cell(row=excel_row, column=excel_col).fill = FIX_FILL


def build_report_workbook(frames: dict[str, DataFrame | None], findings: list,
                          fixes: list[dict], titles: dict[str, str | None] | None = None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    titles = titles or {}
    error_cells = _collect_error_cells(frames, findings)
    fix_cells = _collect_fix_cells(frames, fixes)

    for key in SHEET_KEYS:
        frame = frames.get(key)
        if frame is None:
            continue
        sheet = workbook.create_sheet(title=str(titles.get(key) or key)[:31])
        sheet.append([str(column) for column in frame.columns])
        for _, row in frame.iterrows():
            sheet.append([_cell_value(value) for value in row.tolist()])
        _paint_sheet(sheet, error_cells.get(key, set()), fix_cells.get(key, set()))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
